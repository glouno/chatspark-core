"""File parser registry for local and connector-provided documents."""

from __future__ import annotations

import csv
import mimetypes
import re
import time
from collections.abc import Callable
from dataclasses import dataclass, field
from io import BytesIO, StringIO
from pathlib import Path

from bs4 import BeautifulSoup

from chatspark.documents import ContentType, DocContent, Section
from chatspark.extraction.html import extract_html
from chatspark.extraction.pdf import (
    PdfExtractionError,
    extract_pdf,
    extract_pdf_outcome,
    recognize_ocr,
)
from chatspark.profiles.models import PdfProfile


class UnsupportedDocumentError(ValueError):
    """Raised when no registered parser can safely interpret an input."""


@dataclass(frozen=True)
class ParsedFile:
    content_type: ContentType
    content: DocContent
    mime_type: str
    diagnostics: dict = field(default_factory=dict)


Parser = Callable[[Path, bytes, str], DocContent]

_SUFFIX_TYPES = {
    ".html": ContentType.HTML,
    ".htm": ContentType.HTML,
    ".php": ContentType.HTML,
    ".pdf": ContentType.PDF,
    ".md": ContentType.MARKDOWN,
    ".markdown": ContentType.MARKDOWN,
    ".mdx": ContentType.MARKDOWN,
    ".txt": ContentType.TEXT,
    ".rst": ContentType.TEXT,
    ".docx": ContentType.WORD,
    ".pptx": ContentType.PRESENTATION,
    ".xlsx": ContentType.SPREADSHEET,
    ".csv": ContentType.SPREADSHEET,
    ".tsv": ContentType.SPREADSHEET,
    ".png": ContentType.IMAGE,
    ".jpg": ContentType.IMAGE,
    ".jpeg": ContentType.IMAGE,
    ".tif": ContentType.IMAGE,
    ".tiff": ContentType.IMAGE,
    ".webp": ContentType.IMAGE,
}


def supported_suffixes() -> frozenset[str]:
    return frozenset(_SUFFIX_TYPES)


def detect_file_type(path: Path, mime_type: str | None = None) -> ContentType:
    suffix_type = _SUFFIX_TYPES.get(path.suffix.lower())
    if suffix_type:
        return suffix_type
    mime = (mime_type or "").split(";", 1)[0].strip().lower()
    if mime in {"text/html", "application/xhtml+xml"}:
        return ContentType.HTML
    if mime in {"text/markdown", "text/x-markdown"}:
        return ContentType.MARKDOWN
    if mime == "application/pdf":
        return ContentType.PDF
    if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return ContentType.WORD
    if mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        return ContentType.PRESENTATION
    if mime in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "text/csv",
        "text/tab-separated-values",
    }:
        return ContentType.SPREADSHEET
    if mime.startswith("image/"):
        return ContentType.IMAGE
    if mime.startswith("text/"):
        return ContentType.TEXT
    return ContentType.UNKNOWN


def parse_file(
    path: Path,
    *,
    source_url: str | None = None,
    mime_type: str | None = None,
    pdf_config=None,
    ocr_config=None,
    html_config=None,
) -> ParsedFile:
    path = Path(path)
    raw = read_input(path, pdf_config=pdf_config, mime_type=mime_type)
    resolved_mime = mime_type or mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    return parse_bytes(
        raw,
        name=path.name,
        source_url=source_url or path.as_uri(),
        mime_type=resolved_mime,
        local_path=path,
        pdf_config=pdf_config,
        ocr_config=ocr_config,
        html_config=html_config,
    )


def read_input(path, *, pdf_config=None, mime_type=None):
    """Read PDF/image input with a byte cap, including files that grow after stat."""
    path = Path(path)
    if detect_file_type(path, mime_type) not in {ContentType.PDF, ContentType.IMAGE}:
        return path.read_bytes()
    maximum = (pdf_config or PdfProfile()).max_bytes
    if path.stat().st_size > maximum:
        raise PdfExtractionError("oversized", [])
    with path.open("rb") as stream:
        raw = stream.read(maximum + 1)
    if len(raw) > maximum:
        raise PdfExtractionError("oversized", [])
    return raw


def parse_bytes(
    raw: bytes,
    *,
    name: str,
    source_url: str,
    mime_type: str | None = None,
    local_path: Path | None = None,
    pdf_config=None,
    ocr_config=None,
    html_config=None,
) -> ParsedFile:
    """Parse downloaded bytes while using the source name for format routing."""
    path = local_path or Path(name)
    detected = detect_file_type(Path(name), mime_type)
    resolved_mime = mime_type or mimetypes.guess_type(name)[0] or "application/octet-stream"
    parser = _PARSERS.get(detected)
    if parser is None:
        raise UnsupportedDocumentError(f"Unsupported document type: {path.name} ({resolved_mime})")
    if detected == ContentType.HTML:
        from chatspark.plugins.contracts import ProviderSelection
        from chatspark.runtime.registry import get_registry

        extractor = get_registry().resolve(
            html_config or ProviderSelection(provider="html"), "html"
        )
        content = extractor.extract(raw.decode("utf-8", errors="replace"), source_url)
        return ParsedFile(detected, content, resolved_mime)
    if detected == ContentType.PDF:
        started = time.monotonic()
        outcome = extract_pdf_outcome(raw, pdf_config)
        selective_ocr = (
            ocr_config
            and ocr_config.provider == "tesseract"
            and any(
                len(page.text.strip()) < ocr_config.settings.get("min_native_text_chars", 40)
                for page in outcome.pages
            )
        )
        if ocr_config and (outcome.status == "needs_ocr" or selective_ocr):
            limits = pdf_config or PdfProfile()
            extraction_attempts = outcome.diagnostics.get("attempts", [])
            outcome = recognize_ocr(
                raw,
                ocr_config,
                max_bytes=limits.max_bytes,
                max_output_bytes=limits.max_output_bytes,
                timeout_seconds=limits.timeout_seconds - (time.monotonic() - started),
                native_pages=outcome.pages,
                inspector=limits.inspector,
            )
            outcome.diagnostics["attempts"] = [
                *extraction_attempts,
                *outcome.diagnostics["attempts"],
            ]
        return ParsedFile(
            detected,
            outcome.content,
            resolved_mime,
            {"status": outcome.status, **outcome.diagnostics},
        )
    if detected == ContentType.IMAGE and ocr_config:
        limits = pdf_config or PdfProfile()
        outcome = recognize_ocr(
            raw,
            ocr_config,
            max_bytes=limits.max_bytes,
            max_output_bytes=limits.max_output_bytes,
            timeout_seconds=limits.timeout_seconds,
        )
        return ParsedFile(
            detected,
            outcome.content,
            resolved_mime,
            {"status": outcome.status, **outcome.diagnostics},
        )
    return ParsedFile(detected, parser(path, raw, source_url), resolved_mime)


def _parse_html(path: Path, raw: bytes, source_url: str) -> DocContent:
    return extract_html(raw.decode("utf-8", errors="replace"), source_url)


def _parse_pdf(path: Path, raw: bytes, source_url: str) -> DocContent:
    return extract_pdf(raw)


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _parse_markdown(path: Path, raw: bytes, source_url: str) -> DocContent:
    markdown = _decode_text(raw).replace("\r\n", "\n").strip()
    body = _strip_front_matter(markdown)
    sections = _markdown_sections(body)
    text = _markdown_to_text(body)
    return DocContent(markdown=markdown, text=text, sections=sections or _single_section(text))


def _parse_text(path: Path, raw: bytes, source_url: str) -> DocContent:
    text = _decode_text(raw).replace("\r\n", "\n").strip()
    return DocContent(markdown=text, text=text, sections=_single_section(text))


def _parse_docx(path: Path, raw: bytes, source_url: str) -> DocContent:
    from docx import Document as WordDocument

    document = WordDocument(BytesIO(raw))
    blocks: list[tuple[str | None, str]] = []
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if not value:
            continue
        style = (paragraph.style.name if paragraph.style else "").lower()
        blocks.append((value if style.startswith("heading") else None, value))
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        rendered = _markdown_table(rows)
        if rendered:
            blocks.append((None, rendered))
    return _content_from_blocks(blocks)


def _parse_pptx(path: Path, raw: bytes, source_url: str) -> DocContent:
    from pptx import Presentation

    presentation = Presentation(BytesIO(raw))
    blocks: list[tuple[str | None, str]] = []
    for number, slide in enumerate(presentation.slides, 1):
        title = (slide.shapes.title.text.strip() if slide.shapes.title else "") or f"Slide {number}"
        parts: list[str] = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_text_frame", False):
                value = shape.text.strip()
                if value:
                    parts.append(value)
            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                rendered = _markdown_table(rows)
                if rendered:
                    parts.append(rendered)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if notes:
            parts.append(f"Speaker notes:\n{notes}")
        blocks.append((title, "\n\n".join(parts) or title))
    return _content_from_blocks(blocks)


def _parse_spreadsheet(path: Path, raw: bytes, source_url: str) -> DocContent:
    if path.suffix.lower() in {".csv", ".tsv"}:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        rows = list(csv.reader(StringIO(_decode_text(raw)), delimiter=delimiter))
        return _content_from_blocks([(path.stem, _markdown_table(rows))])

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        blocks = []
        for sheet in workbook.worksheets:
            rows = [
                ["" if value is None else str(value) for value in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if any(cell.strip() for cell in row)]
            blocks.append((sheet.title, _markdown_table(rows)))
        return _content_from_blocks(blocks)
    finally:
        workbook.close()


def _parse_image(path: Path, raw: bytes, source_url: str) -> DocContent:
    # Image evidence requires an explicitly configured OCR provider.
    raise UnsupportedDocumentError("Image input requires an explicit OCR provider")


def _strip_front_matter(markdown: str) -> str:
    if not markdown.startswith("---\n"):
        return markdown
    end = markdown.find("\n---\n", 4)
    return markdown[end + 5 :] if end >= 0 else markdown


def _markdown_to_text(markdown: str) -> str:
    text = re.sub(r"```[^\n]*\n(.*?)```", r"\1", markdown, flags=re.DOTALL)
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"^[>\-*+]\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"[*_~`]", "", text)
    return re.sub(r"\n{3,}", "\n\n", BeautifulSoup(text, "html.parser").get_text("\n")).strip()


def _markdown_sections(markdown: str) -> list[Section]:
    matches = list(re.finditer(r"^(#{1,6})\s+(.+?)\s*$", markdown, re.MULTILINE))
    if not matches:
        return []
    headings: list[str] = []
    sections: list[Section] = []
    for index, match in enumerate(matches):
        level = len(match.group(1))
        headings[level - 1 :] = [match.group(2).strip()]
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(markdown)
        content = _markdown_to_text(markdown[start:end])
        if content:
            sections.append(
                Section(
                    heading_path=list(headings),
                    content=content,
                    start_char_idx=start,
                    end_char_idx=end,
                )
            )
    return sections


def _single_section(text: str) -> list[Section]:
    return (
        [Section(heading_path=[], content=text, start_char_idx=0, end_char_idx=len(text))]
        if text
        else []
    )


def _content_from_blocks(blocks: list[tuple[str | None, str]]) -> DocContent:
    markdown_parts: list[str] = []
    sections: list[Section] = []
    offset = 0
    for heading, value in blocks:
        value = value.strip()
        if not value:
            continue
        rendered = f"## {heading}\n\n{value}" if heading else value
        markdown_parts.append(rendered)
        start = offset + (len(f"## {heading}\n\n") if heading else 0)
        sections.append(
            Section(
                heading_path=[heading] if heading else [],
                content=_markdown_to_text(value),
                start_char_idx=start,
                end_char_idx=start + len(value),
            )
        )
        offset += len(rendered) + 2
    markdown = "\n\n".join(markdown_parts)
    return DocContent(markdown=markdown, text=_markdown_to_text(markdown), sections=sections)


def _markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [
        [str(cell).replace("|", "\\|").replace("\n", " ") for cell in row]
        + [""] * (width - len(row))
        for row in rows
    ]
    header = normalized[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return "\n".join(lines)


_PARSERS: dict[ContentType, Parser] = {
    ContentType.HTML: _parse_html,
    ContentType.PDF: _parse_pdf,
    ContentType.MARKDOWN: _parse_markdown,
    ContentType.TEXT: _parse_text,
    ContentType.WORD: _parse_docx,
    ContentType.PRESENTATION: _parse_pptx,
    ContentType.SPREADSHEET: _parse_spreadsheet,
    ContentType.IMAGE: _parse_image,
}
